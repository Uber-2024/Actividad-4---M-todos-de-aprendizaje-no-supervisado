"""Servidor web para probar el proyecto de aprendizaje no supervisado.

Este módulo sirve una interfaz gráfica desde la carpeta web/ y expone rutas
para generar un dataset sintético, comparar algoritmos de clustering,
consultar métricas y exportar resultados como un libro de Excel.
"""

from __future__ import annotations

import json
from io import BytesIO
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from urllib.parse import urlparse
import sys

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

# Directorios base del proyecto, la interfaz y el paquete de clustering.
ROOT_DIR = Path(__file__).resolve().parent
WEB_DIR = ROOT_DIR / "web"
APP_DIR = ROOT_DIR / "aprendizaje_no_supervisado"

# Se incorpora la raíz del proyecto para poder importar los módulos de Python.
sys.path.insert(0, str(ROOT_DIR))

from aprendizaje_no_supervisado.generate_dataset import generate_dataset
from aprendizaje_no_supervisado.cluster_model import cluster_and_evaluate


class ProjectHandler(BaseHTTPRequestHandler):
    """Maneja las peticiones HTTP del navegador y sirve la interfaz gráfica."""

    def do_GET(self) -> None:
        """Responde con la página principal o archivos estáticos del frontend."""
        parsed = urlparse(self.path)
        route = parsed.path

        # Página de inicio de la aplicación.
        if route in ("/", "/index.html"):
            self._serve_file(WEB_DIR / "index.html", "text/html; charset=utf-8")
            return

        # Devuelve el último reporte de métricas guardado en disco.
        if route == "/api/metrics":
            self._send_json(self._read_metrics())
            return

        # Sirve CSS, JavaScript u otros archivos estáticos de la carpeta web/.
        if route.startswith("/"):
            relative_path = route.lstrip("/")
            file_path = WEB_DIR / relative_path
            if file_path.exists() and file_path.is_file():
                content_type = self._content_type_for(file_path)
                self._serve_file(file_path, content_type)
                return

        self._send_json({"error": "Ruta no encontrada."}, status=404)

    def do_POST(self) -> None:
        """Procesa generación de datos, clustering y exportación a Excel."""
        parsed = urlparse(self.path)
        route = parsed.path

        # Genera el dataset sintético con el tamaño y la semilla indicados.
        if route == "/api/generate-dataset":
            payload = self._read_json_body()
            sample_size = int(payload.get("sample_size", 1200))
            seed = int(payload.get("seed", 2026))
            dataset_path = generate_dataset(sample_size=sample_size, seed=seed)
            self._send_json(
                {
                    "status": "ok",
                    "message": "Dataset generado correctamente.",
                    "sample_size": sample_size,
                    "seed": seed,
                    "path": str(dataset_path.relative_to(ROOT_DIR)),
                }
            )
            return

        # Ejecuta los tres algoritmos de clustering y devuelve las métricas.
        if route == "/api/cluster-model":
            metrics = cluster_and_evaluate()
            self._send_json(
                {
                    "status": "ok",
                    "message": "Agrupamiento finalizado.",
                    "results": metrics,
                    "metrics_path": str(
                        (APP_DIR / "artifacts" / "metrics.json").relative_to(ROOT_DIR)
                    ),
                }
            )
            return

        # Construye y descarga un Excel con la comparación de métodos.
        if route == "/api/export-comparison":
            payload = self._read_json_body()
            # Acepta tanto la respuesta del clustering como el reporte persistido.
            metrics = payload.get("results") or payload.get("models") or payload
            workbook = self._create_comparison_workbook(metrics)
            if workbook is None:
                self._send_json(
                    {"error": "No hay métricas válidas para exportar."}, status=400
                )
                return
            self._send_file_content(
                workbook,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                "comparacion_clustering.xlsx",
            )
            return

        self._send_json({"error": "Endpoint no disponible."}, status=404)

    def log_message(self, format: str, *args: object) -> None:
        """Suprime el ruido habitual de los logs de la librería HTTP."""
        return

    def _read_json_body(self) -> dict:
        """Lee y parsea el contenido JSON enviado por el navegador."""
        content_length = int(self.headers.get("Content-Length", "0"))
        raw_data = self.rfile.read(content_length)
        if not raw_data:
            return {}
        return json.loads(raw_data.decode("utf-8"))

    def _read_metrics(self) -> dict:
        """Carga el archivo de métricas generado por el clustering."""
        metrics_path = APP_DIR / "artifacts" / "metrics.json"
        if not metrics_path.exists():
            return {
                "status": "no-data",
                "message": "Todavía no hay métricas disponibles.",
            }
        return json.loads(metrics_path.read_text(encoding="utf-8"))

    def _serve_file(self, file_path: Path, content_type: str) -> None:
        """Sirve un archivo del disco como respuesta HTTP."""
        content = file_path.read_bytes()
        self.send_response(200)
        self.send_header("Content-Type", content_type)
        self.send_header("Content-Length", str(len(content)))
        self.end_headers()
        self.wfile.write(content)

    def _send_json(self, payload: dict, status: int = 200) -> None:
        """Devuelve una respuesta JSON con el código apropiado."""
        body = json.dumps(payload, ensure_ascii=False, indent=2).encode("utf-8")
        self.send_response(status)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Content-Length", str(len(body)))
        self.end_headers()
        self.wfile.write(body)

    def _send_file_content(self, content: bytes, content_type: str, filename: str) -> None:
        """Envía contenido binario al navegador como un archivo descargable."""
        self.send_response(200)
        self.send_header("Content-Type", content_type)
        self.send_header("Content-Disposition", f'attachment; filename="{filename}"')
        self.send_header("Content-Length", str(len(content)))
        self.end_headers()
        self.wfile.write(content)

    @staticmethod
    def _create_comparison_workbook(metrics: dict) -> bytes | None:
        """Crea un libro Excel con las métricas y resalta la mejor silueta.

        Args:
            metrics: diccionario con resultados por método (kmeans, agglomerative, dbscan).

        Returns:
            Bytes del archivo .xlsx, o None si no hay métricas válidas.
        """
        # Traduce los nombres internos de los métodos a etiquetas legibles en Excel.
        model_names = {
            "kmeans": "K-Means",
            "agglomerative": "Aglomerativo",
            "dbscan": "DBSCAN",
        }
        valid_models = {
            name: metrics[name]
            for name in model_names
            if isinstance(metrics, dict) and isinstance(metrics.get(name), dict)
        }
        if not valid_models:
            return None

        # Crea una hoja con encabezados destacados y una fila para cada método.
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Comparación"
        sheet.append(
            [
                "Método",
                "Clusters",
                "Ruido",
                "Silueta",
                "Calinski-Harabasz",
                "Davies-Bouldin",
            ]
        )
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="245C58")

        for key, label in model_names.items():
            result = valid_models.get(key)
            if result is None:
                continue
            sheet.append(
                [
                    label,
                    result.get("n_clusters"),
                    result.get("n_noise"),
                    result.get("silhouette"),
                    result.get("calinski_harabasz"),
                    result.get("davies_bouldin"),
                ]
            )

        # Resalta la fila con mayor silueta, que es la métrica principal de comparación.
        silhouette_rows = [
            (row, sheet.cell(row=row, column=4).value)
            for row in range(2, sheet.max_row + 1)
            if isinstance(sheet.cell(row=row, column=4).value, (int, float))
        ]
        if silhouette_rows:
            best_row = max(silhouette_rows, key=lambda item: item[1])[0]
            for cell in sheet[best_row]:
                cell.fill = PatternFill("solid", fgColor="D9EAD3")

        for column, width in {
            "A": 18,
            "B": 12,
            "C": 10,
            "D": 12,
            "E": 18,
            "F": 16,
        }.items():
            sheet.column_dimensions[column].width = width
        sheet.freeze_panes = "A2"

        # Serializa el libro en memoria para enviarlo directamente al navegador.
        output = BytesIO()
        workbook.save(output)
        return output.getvalue()

    @staticmethod
    def _content_type_for(file_path: Path) -> str:
        """Determina el tipo MIME real según la extensión del archivo."""
        mime_type, _ = __import__("mimetypes").guess_type(str(file_path))
        if mime_type is None:
            return "application/octet-stream"
        return mime_type


if __name__ == "__main__":
    host = "127.0.0.1"
    # Puerto distinto al de la actividad 3 para poder tener ambos proyectos activos.
    port = 8002
    server = ThreadingHTTPServer((host, port), ProjectHandler)
    print(f"Servidor web activo en http://{host}:{port}")
    try:
        server.serve_forever()
    except KeyboardInterrupt:
        print("\nServidor detenido.")
        server.server_close()
